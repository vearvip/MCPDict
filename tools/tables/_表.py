#!/usr/bin/env python3

from tables import *
import os, re
import logging
from time import time
from collections import defaultdict
from glob import glob
import inspect
from openpyxl import load_workbook
from xlrd import open_workbook
from docx import Document
from docx.enum.text import WD_UNDERLINE
import regex

logging.basicConfig(format='%(message)s', level=logging.INFO)

SOURCE = "data"
TARGET = "output"
VARIANT_FILE = f"tables/{SOURCE}/正字.tsv"

YDS = {"+":"又", "-":"白", "*":"俗", "/":"書","\\":"語","=":"文","?":"存疑", "@": "訓"}
def getYD(py):
	return YDS.get(py[-1], "")

def getCompatibilityVariants():
	d = dict()
	for line in open("../app/src/main/res/raw/orthography_hz_compatibility.txt",encoding="U8"):
		hz, val = line.rstrip()
		d[hz] = val
	return d

def getSTVariants(level=2):
	d = dict()
	for line in open(VARIANT_FILE,encoding="U8"):
		if line.startswith("#"): continue
		fs = line.strip().split("\t")
		if level == 1 and "#" in line:
			continue
		fs[1] = fs[1].split("#")[0].strip()
		if " " not in fs[1]:
			d[fs[0]] = fs[1]
	return d

def getTsvName(xls):
	return re.sub(r" ?(\(\d{0,3}\))+$", "", xls.rsplit(".", 1)[0]) + ".tsv"

def isXlsx(fname):
	return fname.endswith("xlsx")

def isXls(fname):
	return fname.endswith("xls") or fname.endswith("xlsx")

def processFs(v):
	t = type(v)
	if t is float or t is int: return "%d" % v
	if v is None: return ""
	return str(v).strip().replace("\t", " ").replace("\n", " ")

def getXlsxLines(xls, page=0):
	wb = load_workbook(xls, data_only=True)
	sheet = wb.worksheets[page]
	lines = list()
	for row in sheet.rows:
		fs = [processFs(j.value) for j in row[:50]]
		if any(fs):
			line = "\t".join(fs) + "\n"
			lines.append(line)
	return lines

def getXlsLines(xls, page=0):
	wb = open_workbook(xls)
	sheet = wb.sheet_by_index(page)
	lines = list()
	for i in range(sheet.nrows):
		fs = sheet.row_values(i)
		fs = [processFs(j) for j in fs]
		if any(fs):
			line = "\t".join(fs) + "\n"
			lines.append(line)
	return lines

def xls2tsv(xls, page=0):
	tsv = getTsvName(xls)
	if not os.path.exists(xls): return
	if os.path.exists(tsv):
		xtime = os.path.getmtime(xls)
		ttime = os.path.getmtime(tsv)
		if ttime >= xtime: return
	lines = getXlsxLines(xls, page) if isXlsx(xls) else getXlsLines(xls, page)
	t = open(tsv, "w", encoding="U8", newline="\n")
	t.writelines(lines)
	t.close()

def isDocx(fname):
	return fname.endswith("docx")
	
def docx2tsv(doc):
	tsv = getTsvName(doc)
	if not os.path.exists(doc): return
	if os.path.exists(tsv):
		xtime = os.path.getmtime(doc)
		ttime = os.path.getmtime(tsv)
		if ttime >= xtime: return
	lines = []
	for line in Document(doc).paragraphs:
		for run in line.runs:
			if run.font.underline == WD_UNDERLINE.SINGLE:
				for i in run.text:
					lines.append(i + "-")
			elif run.font.underline == WD_UNDERLINE.DOUBLE:
				for i in run.text:
					lines.append(i + "=")
			elif run.font.subscript:
				lines.append(f"{{{run.text}}}")
			else:
				lines.append(run.text)
		lines.append("\n")
	t = open(tsv, "w", encoding="U8", newline="\n")
	t.writelines(lines)
	t.close()

class 表:
	path = os.path.dirname(os.path.abspath(__file__))
	_time = os.path.getmtime(__file__)
	_file = None
	_files = None
	_sep = None
	color = "#1E90FF"
	full = ""
	short = ""
	note = ""
	site = ""
	url = ""
	dictionary = False

	disorder = False
	patches = None
	ybTrimSpace = True
	kCompatibilityVariants = getCompatibilityVariants()
	simplified = 1
	normVariants = getSTVariants(1)
	stVariants = getSTVariants(2)
	isYb = True
	syds = defaultdict(set)
	d = defaultdict(list)
	__mod = None

	def setmod(self, mod):
		self.__mod = mod

	def __str__(self):
		if self.__mod: return self.__mod
		return self.__module__.split(".")[-1]

	def find(self, name):
		if g := glob(name):
			return g
		name = re.sub(".([^.]+)$", "([0-9]).\\1", name)
		return glob(name)

	@property
	def spath(self):
		if self._files:
			self._files = [self.get_fullname(f) for f in self._files]
			self._file = self._files[0]
		sname = self._file
		if not self.short: self.short = self.info["簡稱"]
		if not self.short: self.short = str(self)
		if not sname: sname = f"{self.short}.tsv"
		if not sname.startswith("/"):
			sname = self.get_fullname(sname)
		g = self.find(sname)
		if not g or len(g) != 1:
			if isXls(sname) or isDocx(sname):
				self._file = getTsvName(self._file)
				sname = self.get_fullname(self._file)
				g = self.find(sname)
				if not g or len(g) != 1:
					logging.error(f"\t\t\t{sname} {g}")
					return
			else:
				logging.error(f"\t\t\t未找到{sname} {g}")
				return
		sname = g[0]
		self._file = os.path.basename(sname)
		if isXls(sname):
			page = 1 if self.short in ("中山石岐", "通城") else 0
			if self.short == "開平護龍": page = 3
			xls2tsv(sname, page)
			sname = getTsvName(sname)
		elif isDocx(sname):
			docx2tsv(sname)
			sname = getTsvName(sname)
		return sname

	def get_fullname(self, name):
		return os.path.join(self.path, SOURCE, name)

	@property
	def tpath(self):
		tpath = os.path.join(self.path, TARGET, str(self))
		if not tpath.endswith(".tsv"): tpath += ".tsv"
		return tpath

	def normS(self, s, rep="[\\1]"):
		s = s.replace("(", "（").replace(")", "）")
		s = regex.sub("（((?>[^（）]+|(?R))*)）", rep, s)
		return s

	def normM(self, s, rep="〚\\1〛"):
		s = s.replace("[", "［").replace("]", "］")
		s = regex.sub("［((?>[^［］]+|(?R))*)］", rep, s)
		return s

	def normG(self, s, rep="｛\\1｝"):
		s = s.replace("｛", "{").replace("｝", "}")
		s = regex.sub(r"\{((?>[^\{\}]+|(?R))*)\}", rep, s)
		return s

	def outdated(self):
		classfile = inspect.getfile(self.__class__)
		classtime = os.path.getmtime(classfile)
		varianttime = os.path.getmtime(VARIANT_FILE)
		if classtime < varianttime:
			classtime = varianttime
		spath = self.spath
		if not spath or not os.path.exists(spath):
			return False
		if os.path.exists(self.tpath):
			ftime = os.path.getmtime(spath)
			ttime = os.path.getmtime(self.tpath)
			if ttime < self._time: return True
			if ttime < classtime: return True
			return ttime < ftime
		return True

	def patch(self, d):
		if not self.patches: return
		for hz, py in self.patches.items():
			if not py:
				del d[hz]
				continue
			d[hz] = py.split(",")

	def normAll(self, yb):
		yb = yb.replace("᷉", "̃").replace("ⱼ", "ᶽ")\
			.replace("ʦ", "ts").replace("ʨ", "tɕ").replace("ʧ", "tʃ")\
			.replace("ʣ", "dz").replace("ʥ", "dʑ")\
			.replace("", "ᵑ").replace("", "ᶽ")
		return yb

	def normYb(self, yb):
		if self.isLang() and self.isYb:
			yb = yb.strip()
			yb = yb.replace("Ǿ", "Ǿ").replace("Ǿ", "").lstrip("0∅Ø〇零")
			yb = yb.lower().replace("g", "ɡ").replace("ʼ", "ʰ")
			if not yb.startswith("h") and "h" in yb:
				yb = yb.replace("h", "ʰ")
			if self.ybTrimSpace:
				yb = yb.replace(" ", "")
		return yb

	def isDialect(self):
		return self.langType and not self.langType.startswith("歷史音")

	def isDictionary(self):
		return self.dictionary

	def normJS(self, js):
		if not js: return ""
		last = ""
		l = list()
		for i in js:
			if isHZ(i):
				if last: l.append(last)
				last = ""
				l.append(i)
			else:
				last += i
		if last: l.append(last)
		return " ".join(l)

	def write(self, d):
		self.patch(d)
		t = open(self.tpath, "w", encoding="U8", newline="\n")
		print(f"#漢字\t音標\t解釋", file=t)
		for hz in sorted(d.keys()):
			pys = d[hz]
			hz = self.kCompatibilityVariants.get(hz, hz)
			if self.isDialect() and self.simplified:
				if self.simplified == 1:
					hz = self.normVariants.get(hz, hz)
				else:
					hz = self.stVariants.get(hz, hz)
			if not isHZ(hz):
				if self.isDialect():
					print(f"\t\t\t【{hz}】不是漢字，讀音爲：", ",".join([i.strip() for i in pys]))
				continue
			if self.disorder:
				pys = sorted(pys,key=lambda x:x.split("\t", 1)[0][-1])
			for py in pys:
				if "\t" in py:
					yb, js = py.split("\t", 1)
					js = js.strip().replace("~", "～")
				else:
					yb, js = py, ""
				yb = self.normYb(yb)
				yb = f"{yb}\t{js}"
				yb = self.normAll(yb)
				print(f"{hz}\t{yb}", file=t)
		t.close()

	@property
	def langType(self):
		return self.info["地圖集二分區"]

	def isLang(self):
		return self.langType != None

	@property
	def count(self):
		return len(self.d) + self.unknownCount - (1 if self.unknownCount > 0 else 0)
	
	@property
	def unknownCount(self):
		return len(self.d.get("□", []))

	@property
	def sydCount(self):
		return len(self.syds)

	@property
	def syCount(self):
		return len(set(map(lambda x:x.split("/")[0].rstrip("1234567890"), self.syds.keys())))

	def read(self):
		start = time()
		if self.outdated(): self.update()
		self.syds.clear()
		self.d.clear()
		if not self.tpath or not os.path.exists(self.tpath): return
		for line in open(self.tpath,encoding="U8"):
			line = line.strip()
			if line.startswith("#"): continue
			if "\t" not in line: continue
			hz, py = line.split("\t", 1)
			if self.isLang():
				js = ""
				if "\t" in py: py, js = py.split("\t", 1)
				if js and self.isLang():
					js = self.normJS(js)
				try:
					yd = getYD(py)
				except:
					print("\t\t\t", self.short, py, js)
					exit(1)
				if yd and py.count("*") <= 1:
					js = f"({yd}){js}"
					py = py[:-1]
				if re.match(r"^\([^()]*?\)$", js):
					js = js[1:-1]
				syd = re.sub(r"\(.*?\)","",py).strip(" *|")
				if "-" not in syd:
					self.syds[syd].add(hz)
				if js:
					py += "{%s}" % js
			else:
				if self.isDictionary():
					sep = "▲" if str(self) == "匯纂" else "\t"
					py2, js = py.split(sep, 1)
					py = ("\n\n" if self.d[hz] else "") + py2 + sep + self.normJS(js)
				py = py.replace("\t", "\n")
			if py not in self.d[hz]:
				self.d[hz].append(py)
		passed = time() - start
		logging.info(f"({self.count:5d}({self.unknownCount})-{self.sydCount:4d}-{self.syCount:4d}) {passed:6.3f} {self}")
	
	def load(self, dicts):
		self.read()
		if not self.d: return
		for hz, ybs in self.d.items():
			if hz not in dicts:
				dicts[hz] = {"漢字": hz}
			dicts[hz][str(self)] = "\t".join(ybs)
	
	def parse(self, fs):
		return tuple(fs[:3])

	def format(self, line):
		return line
	
	@property
	def sep(self):
		if self._sep: return self._sep
		sep = "\t"
		spath = self.spath
		if spath.endswith(".csv"): sep = ","
		elif spath.endswith(".tsv"): sep = "\t"
		elif spath.endswith(".txt"): sep = " "
		return sep

	def update(self):
		d = defaultdict(list)
		sep = self.sep
		skip = self.info.get("跳過行數", 0)
		lineno = 0
		files = self._files if self._files else [self.spath]
		for spath in files:
			for line in open(spath,encoding="U8"):
				lineno += 1
				if lineno <= skip: continue
				if line.startswith('#') or line.startswith('"#') : continue
				line = self.format(line)
				fs = [i.strip('" \t') for i in line.strip('\n').split(sep)]
				entries = self.parse(fs)
				if not entries: continue
				if type(entries) is tuple: entries = [entries]
				for fs in entries:
					if len(fs) <= 1: continue
					if len(fs) >= 2:
						hz, yb = fs[:2]
						js = "\t".join(fs[2:])
					if not hz or len(hz) != 1: continue
					if not yb: continue
					if self.isDialect() and isHZ(yb[0]): continue
					p = f"{yb}\t{js}"
					p = p.strip()
					if p not in d[hz]:
						d[hz].append(p)
		self.write(d)

	def splitSySd(self, syd):
		if not syd: return "",""
		tonesymbol = "⁰¹²³⁴⁵⁶"
		tonemark = "˩˨˧˦˥"
		for i in tonesymbol:
			syd = syd.replace(i, str(tonesymbol.index(i)))
		for i in tonemark:
			syd = syd.replace(i, str(tonemark.index(i)+1))
		sy = syd.rstrip("0123456789")
		sd = syd[len(sy):]
		return sy,sd

	def dz2dl(self, sy, dz=None):
		sy = sy.strip()
		if dz is None:
			if "/" in sy:
				return "/".join(map(self.dz2dl, sy.split("/")))
			sy,dz = self.splitSySd(sy)
		if not dz: return sy
		dl = ""
		if dz not in self.toneMaps:
			if dz == "0":
				dl = dz
			elif len(dz) == 1:
				dz = dz + dz
				if dz in self.toneMaps:
					dl = self.toneMaps[dz]
			else:
				dl = "?"
		else:
			dl = self.toneMaps[dz]
		if sy and sy[-1] in "ptkʔ̚" and dz + "0" in self.toneMaps:
			dl = self.toneMaps[dz + "0"]
		return sy + dl
